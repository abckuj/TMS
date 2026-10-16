"""Funnel chart from the recalculated Uber_TAM_SAM_SOM.xlsx (reads cached values)."""
import openpyxl, matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

wb = openpyxl.load_workbook("outputs/Uber_TAM_SAM_SOM.xlsx", data_only=True)
m = wb["Model"]
rows = [(6, "TAM  all-platform India pickup events"),
        (7, "SAM  Uber-serviceable events"),
        (8, "SOM  obtainable in 3-year horizon"),
        (10, "Problem events within SOM"),
        (11, "Events the improvement could address")]
labels = [l for _, l in rows]
cons = [m.cell(r, 4).value for r, _ in rows]
base = [m.cell(r, 5).value for r, _ in rows]
opt = [m.cell(r, 6).value for r, _ in rows]

fig, ax = plt.subplots(figsize=(11, 5.6), dpi=160)
y = range(len(rows))[::-1]
ax.barh(list(y), base, color="#2F6DB5", height=0.5, zorder=2)
for yi, c, b, o in zip(y, cons, base, opt):
    ax.plot([c, o], [yi, yi], color="#1B2A41", lw=2, zorder=3)
    ax.plot([c, c], [yi - .12, yi + .12], color="#1B2A41", lw=2, zorder=3)
    ax.plot([o, o], [yi - .12, yi + .12], color="#1B2A41", lw=2, zorder=3)
    ax.text(o * 1.25, yi, f"base {b:,.1f}M   (range {c:,.2f}–{o:,.0f}M)", va="center", fontsize=9, color="#1B2A41")
ax.set_xscale("log")
ax.set_xlim(0.05, 3e6)
ax.set_yticks(list(y)); ax.set_yticklabels(labels, fontsize=10)
ax.set_xlabel("Annual pickup events, millions (LOG scale)", fontsize=10)
ax.grid(axis="x", color="#DDDDDD", zorder=0)
for s in ("top", "right"): ax.spines[s].set_visible(False)
ax.set_title("Uber Find My Ride — illustrative TAM / SAM / SOM funnel (India, annual)", fontsize=13, loc="left", fontweight="bold")
fig.text(0.01, 0.01,
         "ILLUSTRATIVE MODEL, NOT A VALIDATED MARKET ESTIMATE. Bars = Base case; line = Conservative to Optimistic. Only the 11.6 bn km (Uber, 2025) and market-share inputs are sourced;\n"
         "trip conversion, eligibility, rollout, reach, problem rate and effectiveness are 'Illustrative Assumption — Not Research Evidence'. Source: outputs/Uber_TAM_SAM_SOM.xlsx (Model, Assumptions, Source_Registry).",
         fontsize=7.5, color="#444444", va="bottom")
fig.tight_layout(rect=(0, 0.07, 1, 1))
fig.savefig("charts/TAM_SAM_SOM_funnel.png")
print("ok")
